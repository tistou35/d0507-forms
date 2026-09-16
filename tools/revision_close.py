#!/usr/bin/env python3
# ============================================================
# revision_close.py — ปิดงานฉบับใหม่ที่อนุมัติแล้วในรีโป
#
#   python3 tools/revision_close.py            ดูว่ามีอะไรรอปิด และจะเขียนอะไร
#   python3 tools/revision_close.py --write    เขียนจริง
#
# ── อยู่ตรงไหนของทั้งสาย ─────────────────────────────────────
#   1. หน้าเว็บ  แก้ต้นฉบับ → ออกฉบับใหม่ → ฉบับร่าง + DRF          (app.js)
#   2. หน้าอนุมัติ อนุมัติ DRF → ฉบับร่างติดป้าย APPROVED เปิดให้อ่าน     (Revision.gs)
#                  revisions.state = approved · ต้นฉบับที่ใช้อยู่ยังไม่ขยับ
#   3. ตัวนี้     ออกฉบับใหม่จริง แล้วตั้ง state = closed
#
# ทำไมไม่ใช้ฉบับร่างเป็นต้นฉบับใหม่ตรง ๆ: ฉบับร่างคัดลอกมาจากฉบับเดิม เลขกำกับ
# และวันมีผลในตัวมันจึงเป็นของเก่า ทดสอบแล้วได้ฟอร์มเปล่าพิมพ์ QA-DRF-301-A
# ขณะที่ .docx กับทะเบียนเป็น -B แผ่นเดียวอ้างสองเลข ต้นฉบับใหม่จึงสร้างจาก .docx
# หลังเลื่อนเลขแล้วเสมอ
#
# ทำไมข้อ 3 ไม่ทำบนเว็บด้วย: .docx ฉบับจริง ทะเบียน และฟอร์มเปล่า PDF อยู่ในรีโป
# (กฎของโครงการ: .docx คือฉบับจริง) หน้าเว็บที่ host สาธารณะเขียนรีโปไม่ได้
# และไม่ควรเขียนได้
#
# ── สิ่งที่ทำต่อหนึ่งรายการ ──────────────────────────────────
#   · ดาวน์โหลดฉบับร่างที่อนุมัติเป็น .docx ทับของเดิม — ของเดิมเก็บเป็น *-issXXrevYY-archive.docx
#   · แทนวันมีผลและ Issue/Rev · เลื่อนตัวอักษรท้ายเลขกำกับ (bump_code.py)
#   · สร้างต้นฉบับบน Google Docs ใหม่จาก .docx นั้น ตั้งอ่านได้อย่างเดียว
#   · ต้นฉบับเดิม → เปลี่ยนชื่อเป็น superseded ย้ายไป D-0507 Superseded sources
#   · ทะเบียน: iss · rev · eff · edit = id ต้นฉบับใหม่  และ publications.json ถ้ามี
#   · ทำฟอร์มเปล่า PDF ใหม่จากต้นฉบับใหม่ (make_blank.py)
#
# ── สิ่งที่ "ไม่" ทำ — ต้องทำเองและเขียนไว้ท้ายผลลัพธ์ทุกครั้ง ─────
#   · ไม่เปิด PDF ดูแทนคน (กฎ: ตรวจด้วยการเปิดไฟล์ ไม่ใช่อ่าน log)
#   · ไม่สร้างแม่แบบกรอก (tpl) ใหม่ — importTemplate ทิ้งการวาง token ด้วยมือ
#   · ไม่เขียน Revision Log ใน CLAUDE.md — พิมพ์แถวร่างไว้ให้
#   · ไม่ทับ .docx ที่มี {{token}} ฝังอยู่ (เช่น PWR) — ต้นฉบับบน Docs ถอด token
#     ออกไปแล้ว ทับไปจะเสีย token ที่วางไว้ในกระดาษ ต้องรวมเนื้อหาด้วยมือ
# ============================================================
import datetime
import json
import os
import re
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(HERE, 'tools'))
SRC_DIR = os.path.dirname(HERE)
REG = os.path.join(HERE, 'forms_register.json')
PUBS = os.path.join(HERE, 'publications.json')
MON = ['JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC']


def node(*args):
    r = subprocess.run(['node', os.path.join(HERE, 'tools', 'revisions.mjs')] + list(args),
                       capture_output=True, text=True, cwd=HERE)
    if r.returncode:
        sys.exit('revisions.mjs ล้ม: ' + r.stderr.strip()[-400:])
    return r.stdout


def dmmm(iso):
    """2026-10-01 → 01 OCT 2026 · ค่าว่างหรืออ่านไม่ออก → ''"""
    try:
        d = datetime.date.fromisoformat(str(iso)[:10])
        return '%02d %s %d' % (d.day, MON[d.month - 1], d.year)
    except ValueError:
        return ''


def has_tokens(path):
    with zipfile.ZipFile(path) as z:
        return any(b'{{' in z.read(n) for n in z.namelist() if n.endswith('.xml'))


def replace_text(path, pairs):
    """แทนข้อความใน .docx ทั้งตัวเอกสาร หัว และท้ายกระดาษ — คืนจำนวนจุดต่อคู่"""
    import docx
    d = docx.Document(path)
    hits = {old: 0 for old, _ in pairs}

    def pars():
        for p in d.paragraphs:
            yield p
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    yield from c.paragraphs
        for s in d.sections:
            for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                yield from part.paragraphs
                for t in part.tables:
                    for row in t.rows:
                        for c in row.cells:
                            yield from c.paragraphs

    seen = set()
    for p in pars():
        if id(p._p) in seen:          # เซลล์ที่ merge กันคืน paragraph เดิมซ้ำ
            continue
        seen.add(id(p._p))
        full = ''.join(r.text for r in p.runs)
        new = full
        for old, rep in pairs:
            if old and old in new:
                hits[old] += new.count(old)
                new = new.replace(old, rep)
        if new != full and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    d.save(path)
    return hits


SUPERSEDED = 'D-0507 Superseded sources'


def supersede(fid, name, tk):
    """เปลี่ยนชื่อต้นฉบับเดิมแล้วย้ายเข้าโฟลเดอร์ฉบับที่ถูกแทน — ไม่ลบ
    ฉบับที่ถูกแทนต้องเก็บไว้ (กฎของโครงการ) สิทธิ์อ่านเดิมคงอยู่ ลิงก์ที่เคยแจกยังเปิดได้
    คืนข้อความสรุป หรือ '' ถ้าทำไม่ได้ (ไฟล์ที่เครื่องมือนี้ไม่ได้สร้าง drive.file มองไม่เห็น)"""
    import urllib.error
    import urllib.parse
    import make_edit_doc as m
    try:
        q = urllib.parse.quote("mimeType='application/vnd.google-apps.folder' and trashed=false and name='%s'"
                               % SUPERSEDED)
        fs = json.load(m.api('https://www.googleapis.com/drive/v3/files?q=%s&fields=files(id)' % q, tk))['files']
        dest = fs[0]['id'] if fs else json.load(m.api(
            'https://www.googleapis.com/drive/v3/files?fields=id', tk,
            json.dumps({'name': SUPERSEDED, 'mimeType': 'application/vnd.google-apps.folder'}).encode(),
            'application/json'))['id']
        cur = json.load(m.api('https://www.googleapis.com/drive/v3/files/%s?fields=parents' % fid, tk))
        url = ('https://www.googleapis.com/drive/v3/files/%s?addParents=%s&removeParents=%s&fields=id'
               % (fid, dest, ','.join(cur.get('parents', []))))
        m.api(url, tk, json.dumps({'name': name}).encode(), 'application/json', 'PATCH')
        return '→ %s / %s' % (SUPERSEDED, name)
    except urllib.error.HTTPError:
        return ''


def plan(r, reg, pubs):
    row = next((f for f in reg['forms'] if f.get('doc') == r['doc']), None)
    out = {'r': r, 'row': row, 'stop': ''}
    if not row:
        out['stop'] = 'ไม่มี %s ในทะเบียน' % r['doc']
        return out
    if r.get('promoteError'):
        out['stop'] = 'ตอนอนุมัติติดป้าย/เปิดสิทธิ์อ่านฉบับร่างไม่สำเร็จ: %s' % r['promoteError']
        return out
    if row.get('edit') != r.get('liveId'):
        out['stop'] = ('ทะเบียนชี้ต้นฉบับ %s ไม่ใช่ตัวที่ฉบับร่างคัดลอกมา (%s) — มีคนเปลี่ยนต้นฉบับ'
                       'ระหว่างทาง ต้องดูเองก่อน' % (row.get('edit'), r.get('liveId')))
        return out
    if (row.get('iss'), row.get('rev')) != (r.get('iss'), r.get('rev')):
        out['stop'] = ('ทะเบียนเป็น Issue %s / Rev %s แต่ฉบับร่างตั้งต้นจาก Issue %s / Rev %s'
                       % (row.get('iss'), row.get('rev'), r.get('iss'), r.get('rev')))
        return out
    path = os.path.join(SRC_DIR, row.get('docx') or '')
    if not row.get('docx') or not os.path.exists(path):
        out['stop'] = 'ไม่พบ .docx ต้นฉบับ %s' % (row.get('docx') or '(ไม่ระบุ)')
        return out
    if has_tokens(path):
        out['stop'] = ('%s มี {{token}} ฝังอยู่ — ต้นฉบับบน Docs ถอด token ออกไปแล้ว ทับไปจะเสีย '
                       'ต้องรวมเนื้อหาใหม่เข้า .docx ด้วยมือ' % row['docx'])
        return out
    stem = row['docx'][:-5]
    out.update(path=path,
               archive=os.path.join(SRC_DIR, '%s-iss%srev%s-archive.docx' % (stem, r['iss'], r['rev'])),
               eff=dmmm(r.get('newEff')),
               pub=next((p for p in pubs['pubs'] if p.get('doc') == r['doc']), None))
    if not out['eff']:
        out['stop'] = 'DRF ไม่มีวันมีผล (newEff) — ผู้อนุมัติต้องกรอกก่อน'
    elif os.path.exists(out['archive']):
        out['stop'] = 'มีไฟล์ archive ชื่อ %s อยู่แล้ว — ไม่เขียนทับ' % os.path.basename(out['archive'])
    return out


def main():
    write = '--write' in sys.argv
    todo = json.loads(node('list', 'approved'))
    if not todo:
        print('ไม่มีฉบับใหม่ที่อนุมัติแล้วรอปิดงาน')
        return
    reg = json.load(open(REG, encoding='utf-8'))
    pubs = json.load(open(PUBS, encoding='utf-8'))

    for r in todo:
        p = plan(r, reg, pubs)
        row = p['row'] or {}
        print('\n── %s  Issue %s/Rev %s → Issue %s/Rev %s  (%s · DRF %s)' % (
            r['doc'], r.get('iss'), r.get('rev'), r.get('newIss'), r.get('newRev'),
            r['id'], r.get('tracking') or '—'))
        if p['stop']:
            print('   🔴 ข้าม: ' + p['stop'])
            continue
        print('   .docx    %s  (เดิมเก็บเป็น %s)' % (row['docx'], os.path.basename(p['archive'])))
        print('   ฉบับร่าง %s → ต้นฉบับใหม่ (สร้างจาก .docx หลังเลื่อนเลข)' % r['draftId'])
        print('   ต้นฉบับเดิม %s → superseded' % row['edit'])
        print('   วันมีผล  %s → %s' % (row.get('eff'), p['eff']))
        print('   เลขกำกับ %s → เลื่อนตัวอักษรด้วย bump_code.py' % row.get('code'))
        if p['pub']:
            print('   publications.json  %s' % p['pub']['id'])
        if not write:
            continue

        import make_blank
        import bump_code
        data = make_blank.public_export(r['draftId'], 'docx')
        os.rename(p['path'], p['archive'])
        open(p['path'], 'wb').write(data)

        # วันมีผลกับ Issue/Rev — รูปแบบในกระดาษแต่ละใบไม่เหมือนกัน แทนเฉพาะที่ตรงตัว
        # และรายงานจำนวนจุดให้เห็น ไม่เจอเลยไม่ได้แปลว่าผิด (บางใบไม่พิมพ์ Issue/Rev)
        oi, orv, ni, nr = r['iss'], r['rev'], r['newIss'], r['newRev']
        hits = replace_text(p['path'], [
            (row.get('eff') or '', p['eff']),
            ('Issue %s Rev %s' % (oi, orv), 'Issue %s Rev %s' % (ni, nr)),
            ('Issue %s / Rev %s' % (oi, orv), 'Issue %s / Rev %s' % (ni, nr)),
            ('ISSUE NO. %s/REVISION NO. %s' % (oi, orv), 'ISSUE NO. %s/REVISION NO. %s' % (ni, nr)),
        ])
        for k, v in hits.items():
            print('   แทน %-34s %d จุด' % ('"%s"' % k, v))

        row.update(iss=ni, rev=nr, eff=p['eff'])
        json.dump(reg, open(REG, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        open(REG, 'a', encoding='utf-8').write('\n')
        if bump_code.main([row['abbr']]):
            print('   🔴 bump_code รายงานปัญหา — ตรวจเลขกำกับในเอกสารเอง')
        reg = json.load(open(REG, encoding='utf-8'))          # bump_code เขียนทะเบียนใหม่
        row = next(f for f in reg['forms'] if f.get('doc') == r['doc'])

        # ต้นฉบับใหม่บน Docs — จาก .docx ที่เลื่อนเลขแล้ว ไม่ใช่จากฉบับร่าง
        import make_edit_doc
        import share_edit_doc
        from gas_push import token
        tk = token()
        old_edit = row['edit']
        new_edit = make_edit_doc.upload(p['path'], '%s — %s' % (row['doc'], row.get('t') or ''),
                                        make_edit_doc.folder_id(tk), tk)
        _, ok = share_edit_doc.set_reader(new_edit, tk)
        print('   ต้นฉบับใหม่ %s%s' % (new_edit, '' if ok else '  🔴 ตั้งอ่านได้อย่างเดียวไม่สำเร็จ'))
        moved = supersede(old_edit, '%s — Issue %s Rev %s — superseded' % (row['doc'], oi, orv), tk)
        print('   ต้นฉบับเดิม ' + (moved or '🔴 ย้ายไม่ได้ (ไฟล์ไม่ได้สร้างจากรีโปนี้) — เปลี่ยนชื่อ/ย้ายเองในไดรฟ์'))
        row['edit'] = new_edit
        json.dump(reg, open(REG, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        open(REG, 'a', encoding='utf-8').write('\n')

        if p['pub']:
            p['pub'].update(edit=new_edit, ed=ni, rev=nr, code=row['code'],
                            eff=str(r['newEff'])[:10])
            json.dump(pubs, open(PUBS, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

        if row.get('blank'):
            make_blank.main([row['abbr']])
        node('close', r['id'], new_edit)

        print('\n   ✅ ปิดแล้ว — ที่เหลือต้องทำเอง:')
        print('      1. python3 build.py แล้ว "เปิด" %s ดูว่าหัว ท้าย เลขกำกับ วันมีผล ตรงกันทุกที่' % (row.get('blank') or row['docx']))
        if row.get('tpl'):
            print('      2. แม่แบบกรอก (tpl) ยังเป็นฉบับเดิม — ทำใหม่ด้วย importTemplate(\'%s\') '
                  'แล้ววาง token ที่เคยวางมือซ้ำ' % row['abbr'])
        print('      3. Revision Log ใน CLAUDE.md:')
        print('         | %s | %s | %s | <สรุปจาก DRF %s> | <เหตุผล> | %s | %s |' % (
            nr, p['eff'], r['doc'], r.get('tracking') or '', r.get('byName') or '',
            r.get('approvedBy') or ''))
        print('      4. commit')

    if not write:
        print('\n(ดูอย่างเดียว — ใส่ --write เพื่อเขียนจริง)')


if __name__ == '__main__':
    main()
